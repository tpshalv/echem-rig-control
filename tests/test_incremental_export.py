import csv
from datetime import UTC, datetime, timedelta
import json
from pathlib import Path
from threading import Event, Thread

from openpyxl import load_workbook
import pytest

from rig_control.data.export import build_wide_rows, export_excel
from rig_control.data.incremental_export import IncrementalCsvExporter


START = datetime(2026, 9, 24, tzinfo=UTC)


def record(second, value, channel="temperature"):
    return dict(device_id="probe", channel=channel, unit="degC", value=value,
                timestamp=(START + timedelta(seconds=second)).isoformat(), quality="good")


def prepare(directory):
    (directory / "metadata.json").write_text('{"experiment_id": "test"}', encoding="utf-8")
    (directory / "measurements.journal.jsonl").touch()
    (directory / "events.journal.jsonl").touch()
    return IncrementalCsvExporter(directory)


def append(directory, records):
    with (directory / "measurements.journal.jsonl").open("a", encoding="utf-8") as stream:
        for item in records:
            stream.write(json.dumps(item) + "\n")


def rows(directory):
    with (directory / "measurements-wide.csv").open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.reader(stream))


def test_live_csv_appends_without_replacing_old_rows_and_retains_tail(tmp_path, monkeypatch):
    exporter = prepare(tmp_path)
    append(tmp_path, [record(i, 20 + i) for i in range(5)])
    exporter.update()
    before = exporter.path.read_bytes()
    assert len(rows(tmp_path)) == 4  # Header and seconds 0, 1, 2; latest two held.
    append(tmp_path, [record(i, 20 + i) for i in range(5, 10)])

    def no_replacement(*args):
        raise AssertionError("An ordinary live update must append, not replace")

    monkeypatch.setattr(Path, "replace", no_replacement)
    # A new exporter instance also resumes from the persisted byte offset.
    IncrementalCsvExporter(tmp_path).update()
    assert exporter.path.read_bytes().startswith(before)
    assert len(rows(tmp_path)) == 9
    exporter.update(final=True)
    assert len(rows(tmp_path)) == 11
    exporter.update(final=True)
    assert len(rows(tmp_path)) == 11  # No duplicate rows after repeated export.


def test_late_readings_new_channels_and_averages_match_full_export(tmp_path):
    exporter = prepare(tmp_path)
    first = [record(0.1, 20), record(1, 21), record(5, 25)]
    append(tmp_path, first)
    exporter.update()
    later = [record(0.2, 22), record(0.3, 30, "target_setpoint"),
             record(0.4, 35, "target_setpoint"), record(8, 28)]
    append(tmp_path, later)
    exporter.update(final=True)
    headers, expected = build_wide_rows(first + later)
    assert rows(tmp_path) == [headers, *[[str(cell) for cell in row] for row in expected]]


def test_partial_journal_line_is_not_consumed_until_complete(tmp_path):
    exporter = prepare(tmp_path)
    path = tmp_path / "measurements.journal.jsonl"
    encoded = json.dumps(record(0, 21))
    path.write_text(encoded[:20], encoding="utf-8")
    exporter.update(final=True)
    assert rows(tmp_path) == [["timestamp_utc"]]
    with path.open("a", encoding="utf-8") as stream:
        stream.write(encoded[20:] + "\n")
    exporter.update(final=True)
    assert rows(tmp_path)[1][1] == "21.0"


def test_retry_after_failed_csv_write_does_not_double_count(tmp_path, monkeypatch):
    exporter = prepare(tmp_path)
    append(tmp_path, [record(0, 20)])
    exporter.update(final=True)
    append(tmp_path, [record(0.1, 22)])
    replace = Path.replace

    def locked(*args):
        raise PermissionError("CSV is open elsewhere")

    monkeypatch.setattr(Path, "replace", locked)
    with pytest.raises(PermissionError):
        exporter.update(final=True)
    monkeypatch.setattr(Path, "replace", replace)
    exporter.update(final=True)
    assert rows(tmp_path)[1][1] == "21.0"
    # Recover a partially appended CSV (e.g. process interrupted before commit).
    with exporter.path.open("ab") as stream:
        stream.write(b"partial,row")
    exporter.update(final=True)
    assert len(rows(tmp_path)) == 2
    assert rows(tmp_path)[1][1] == "21.0"


def test_excel_has_no_raw_sheet_splits_data_and_preserves_subsecond_interval(tmp_path, monkeypatch):
    prepare(tmp_path)
    append(tmp_path, [record(i / 2, 20 + i) for i in range(8)])
    IncrementalCsvExporter(tmp_path, 0.5).update(final=True)
    monkeypatch.setattr("rig_control.data.export.EXCEL_ROW_LIMIT", 4)
    path = export_excel(tmp_path)  # Infer saved half-second interval after restart.
    workbook = load_workbook(path, read_only=True)
    try:
        assert workbook.sheetnames == ["Overview", "Data", "Data 2", "Data 3", "Events"]
        data = [row for name in ("Data", "Data 2", "Data 3")
                for row in list(workbook[name].values)[1:]]
        assert len(data) == 8
        assert [row[1] for row in data] == list(range(20, 28))
        assert data[1][0].microsecond == 500_000
        assert workbook["Data"]["A2"].number_format.endswith(".000")
    finally:
        workbook.close()


def test_snapshot_stays_consistent_while_new_readings_are_exported(tmp_path):
    exporter = prepare(tmp_path)
    append(tmp_path, [record(0, 20), record(1, 21)])
    exporter.update(final=True)
    with exporter.snapshot() as (_, snapshot_rows):
        append(tmp_path, [record(2, 22)])
        exporter.update(final=True)
        assert len(list(snapshot_rows)) == 2
    assert len(rows(tmp_path)) == 4


def test_on_demand_excel_does_not_block_recording(tmp_path, monkeypatch):
    from rig_control.data.directory_writer import DirectoryExperimentWriter
    from rig_control.data.experiment import ExperimentMetadata
    from rig_control.data.records import MeasurementRecord
    from rig_control.experiment_recording import ExperimentRecorder
    from rig_control.models import Measurement
    from rig_control.polling import PollingBatch

    recorder = ExperimentRecorder(writer_factory=lambda path: DirectoryExperimentWriter(
        path, live_export_interval_seconds=0))
    recorder.start(metadata=ExperimentMetadata("test", "tester"), root_directory=tmp_path)
    entered, release, recorded = Event(), Event(), Event()
    failures = []
    real_export = __import__("rig_control.data.export", fromlist=["export_excel"]).export_excel

    def delayed(*args, **kwargs):
        entered.set()
        assert release.wait(5)
        return real_export(*args, **kwargs)

    monkeypatch.setattr("rig_control.data.directory_writer.export_excel", delayed)

    def exporting():
        try:
            recorder.export_excel()
        except Exception as error:
            failures.append(error)

    thread = Thread(target=exporting)
    thread.start()
    assert entered.wait(5)

    def acquire():
        item = MeasurementRecord("probe", "temperature", Measurement(25, "degC", START))
        recorder.record_batch(PollingBatch(START, START, (item,), (), ()))
        recorded.set()

    acquisition = Thread(target=acquire)
    acquisition.start()
    try:
        assert recorded.wait(2), "Excel export held the acquisition lock"
        assert recorder.is_recording
    finally:
        release.set()
        thread.join(5)
        acquisition.join(5)
    assert not failures
    assert recorder.diagnostic_metrics()["measurement_records_written"] == 1
    recorder.stop()
    assert (recorder.experiment_directory / "experiment.xlsx").exists()


def test_export_rejects_non_experiment_folder_without_creating_cache(tmp_path):
    with pytest.raises(ValueError, match="Choose an experiment folder"):
        export_excel(tmp_path)
    assert not (tmp_path / "export-cache.sqlite3").exists()
