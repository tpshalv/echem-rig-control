from collections import defaultdict
from datetime import datetime, timezone
import json
from math import floor
from pathlib import Path
from typing import Any


LAST_VALUE_CHANNELS = {
    "active_setpoint",
    "target_setpoint",
    "alarm_state",
    "error_status",
}


def export_experiment_files(
    experiment_directory: str | Path,
    *,
    bin_seconds: float = 1.0,
) -> tuple[Path, Path]:
    """Finish the incremental CSV and stream a compact Excel workbook."""
    csv_path = export_wide_csv(experiment_directory, bin_seconds=bin_seconds)
    return csv_path, export_excel(experiment_directory, bin_seconds=bin_seconds,
                                  refresh=False)


def export_wide_csv(
    experiment_directory: str | Path,
    *,
    bin_seconds: float = 1.0,
    final: bool = True,
) -> Path:
    from rig_control.data.incremental_export import IncrementalCsvExporter

    return IncrementalCsvExporter(experiment_directory, bin_seconds).update(final=final)


def export_excel(
    experiment_directory: str | Path,
    *,
    bin_seconds: float | None = None,
    refresh: bool = True,
) -> Path:
    """Export a snapshot without stopping recording or loading the run into RAM."""
    from rig_control.data.incremental_export import IncrementalCsvExporter

    directory = Path(experiment_directory)
    for name in ("metadata.json", "measurements.journal.jsonl", "events.journal.jsonl"):
        if not (directory / name).is_file():
            raise ValueError(f"Choose an experiment folder containing {name}")
    if bin_seconds is None:
        # A folder selected after restart keeps its recorded export interval.
        import sqlite3

        cache = directory / "export-cache.sqlite3"
        bin_seconds = 1.0
        if cache.exists():
            connection = sqlite3.connect(cache)
            try:
                row = connection.execute(
                    "SELECT value FROM state WHERE key='bin_seconds'"
                ).fetchone()
                if row is not None:
                    bin_seconds = float(json.loads(row[0]))
            finally:
                connection.close()
    exporter = IncrementalCsvExporter(directory, bin_seconds)
    if refresh:
        state_path = directory / "recording-state.json"
        recording = (state_path.exists() and json.loads(
            state_path.read_text(encoding="utf-8")).get("state") == "recording")
        exporter.update(final=not recording)
    metadata = json.loads((directory / "metadata.json").read_text(encoding="utf-8"))
    path = directory / "experiment.xlsx"
    # Capture an event boundary as well, so an active journal cannot keep an
    # on-demand export running indefinitely.
    events_path = directory / "events.journal.jsonl"
    events_end = events_path.stat().st_size
    with exporter.snapshot() as (signals, rows):
        _write_xlsx(path, metadata=metadata,
                    wide_headers=exporter.headers(signals, metadata), wide_rows=rows,
                    events=_iter_jsonl(events_path, events_end), bin_seconds=bin_seconds)
    return path


def build_wide_rows(
    measurements: list[dict[str, Any]],
    *,
    bin_seconds: float = 1.0,
    channel_labels: dict[str, dict[str, str]] | None = None,
) -> tuple[list[str], list[list[object]]]:
    """Bin long-form observations onto a regular grid without carrying values."""
    if bin_seconds <= 0:
        raise ValueError("Export bin size must be positive")
    signals: dict[tuple[str, str], str] = {}
    buckets: dict[int, dict[tuple[str, str], list[dict[str, Any]]]] = defaultdict(
        lambda: defaultdict(list)
    )
    for record in measurements:
        key = (str(record["device_id"]), str(record["channel"]))
        unit = str(record["unit"])
        previous_unit = signals.setdefault(key, unit)
        if previous_unit != unit:
            raise ValueError(f"Signal {key} changed unit from {previous_unit} to {unit}")
        timestamp = datetime.fromisoformat(str(record["timestamp"]))
        if timestamp.tzinfo is None:
            raise ValueError("Measurement timestamp has no timezone")
        bucket = floor(timestamp.timestamp() / bin_seconds + 0.5)
        buckets[bucket][key].append(record)

    ordered_signals = sorted(signals)
    headers = ["timestamp_utc"] + [
        f"{device_id}.{channel}"
        + (f" ({channel_labels[device_id][channel]})"
           if channel_labels and channel in channel_labels.get(device_id, {}) else "")
        + f" [{signals[(device_id, channel)]}]"
        for device_id, channel in ordered_signals
    ]
    rows: list[list[object]] = []
    for bucket in sorted(buckets):
        timestamp = datetime.fromtimestamp(
            bucket * bin_seconds, tz=timezone.utc
        ).isoformat()
        row: list[object] = [timestamp]
        for signal in ordered_signals:
            records = buckets[bucket].get(signal, [])
            if not records:
                row.append("")
                continue
            if signal[1] in LAST_VALUE_CHANNELS:
                selected = max(records, key=lambda item: str(item["timestamp"]))
                row.append(float(selected["value"]))
            else:
                values = [float(record["value"]) for record in records]
                row.append(sum(values) / len(values))
        rows.append(row)
    return headers, rows


def _channel_labels(metadata: dict[str, Any]) -> dict[str, dict[str, str]]:
    """Labels are snapshotted at recording start; raw channel IDs never change."""
    labels = json.loads(metadata.get("extra", {}).get("channel_labels_json", "{}"))
    if not isinstance(labels, dict) or any(
        not isinstance(channels, dict)
        or any(not isinstance(label, str) for label in channels.values())
        for channels in labels.values()
    ):
        raise ValueError("Invalid recorded channel labels")
    return labels


def _iter_jsonl(path: Path, end: int):
    with path.open("rb") as stream:
        while stream.tell() < end:
            line = stream.readline(end - stream.tell())
            if not line.endswith(b"\n"):
                break
            if line.strip():
                value = json.loads(line)
                if not isinstance(value, dict):
                    raise ValueError(f"{path.name} contains a non-object record")
                yield value


EXCEL_ROW_LIMIT = 1_048_576
EXCEL_COLUMN_LIMIT = 16_384


def _write_xlsx(
    path: Path,
    *,
    metadata: dict[str, Any],
    wide_headers: list[str],
    wide_rows,
    events,
    bin_seconds: float,
) -> None:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    if len(wide_headers) > EXCEL_COLUMN_LIMIT:
        raise ValueError("Too many signals for an Excel worksheet; use the CSV export")
    workbook = Workbook(write_only=True)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    time_format = "yyyy-mm-dd hh:mm:ss" + (".000" if bin_seconds < 1 else "")

    def cells(sheet, values, *, header=False):
        result = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            # Labels and operator notes must remain text, including leading '='.
            if isinstance(value, str):
                cell.data_type = "s"
            if isinstance(value, datetime):
                cell.number_format = time_format
            if header:
                cell.fill, cell.font = header_fill, header_font
            result.append(cell)
        return result

    def table(name, headers, rows):
        part = 0
        sheet = None
        count = 0
        for row in rows:
            if sheet is None or count >= EXCEL_ROW_LIMIT:
                if sheet is not None:
                    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{count}"
                part += 1
                sheet = workbook.create_sheet(name if part == 1 else f"{name} {part}")
                sheet.freeze_panes = "A2"
                sheet.sheet_view.showGridLines = False
                for index, label in enumerate(headers, 1):
                    sheet.column_dimensions[get_column_letter(index)].width = (
                        23 if index == 1 else min(max(len(label) + 2, 16), 45)
                    )
                sheet.append(cells(sheet, headers, header=True))
                count = 1
            sheet.append(cells(sheet, row))
            count += 1
        if sheet is None:
            sheet = workbook.create_sheet(name)
            sheet.freeze_panes = "A2"
            sheet.append(cells(sheet, headers, header=True))
            count = 1
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{count}"

    overview = workbook.create_sheet("Overview")
    overview.column_dimensions["A"].width = 34
    overview.column_dimensions["B"].width = 85
    overview.freeze_panes = "A4"
    overview.sheet_view.showGridLines = False
    overview.append(cells(overview, ["Experiment export"]))
    overview.append([])
    overview.append(cells(overview, ["Field", "Value"], header=True))
    for key, value in metadata.items():
        if key == "extra":
            for extra_key, extra_value in dict(value).items():
                overview.append(cells(overview, [f"extra.{extra_key}", extra_value]))
        else:
            overview.append(cells(overview, [key, value]))
    overview.append(cells(overview, ["Time interval (seconds)", bin_seconds]))
    overview.append(cells(overview, ["Binning method",
        "Nearest time bin; readings averaged, setpoints use latest; missing signals blank."]))
    overview.append(cells(overview, ["Original readings",
        "Full readings and timestamps remain in measurements.journal.jsonl."]))
    overview.append(cells(overview, ["Exported at (UTC)", datetime.now(timezone.utc).replace(tzinfo=None)]))
    table("Data", wide_headers,
          ([_excel_datetime(row[0]), *row[1:]] for row in wide_rows))
    event_headers = ["timestamp", "severity", "source", "message"]
    table("Events", event_headers,
          ([_excel_datetime(event.get("timestamp")),
            *(event.get(key) for key in event_headers[1:])] for event in events))
    temporary = path.with_name(f"{path.stem}.tmp{path.suffix}")
    try:
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)


def _excel_datetime(value: object) -> datetime | object:
    if not isinstance(value, str):
        return value
    parsed = datetime.fromisoformat(value)
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone(timezone.utc).replace(tzinfo=None)
    return parsed
