import csv
from datetime import UTC, datetime, timedelta
import json

from openpyxl import load_workbook
import pytest

from rig_control.data.annotations import (
    ANNOTATIONS_JOURNAL,
    AnnotationRecord,
    RunReview,
    annotation_record_from_dict,
    annotation_record_to_dict,
    append_run_review,
    deleted_annotation,
    edited_annotation,
    format_annotation_cell,
    latest_run_review,
    new_annotation,
    nice_uncertainty,
    read_annotation_journal,
    read_annotations,
    resolve_annotations,
)
from rig_control.data.directory_writer import DirectoryExperimentWriter
from rig_control.data.experiment import ExperimentMetadata
from rig_control.data.export import build_wide_rows, export_excel
from rig_control.data.in_memory_writer import InMemoryExperimentWriter
from rig_control.data.incremental_export import IncrementalCsvExporter
from rig_control.experiment_recording import ExperimentRecorder


START = datetime(2026, 10, 22, 9, 0, tzinfo=UTC)


def at(seconds: float) -> datetime:
    return START + timedelta(seconds=seconds)


def note(seconds, text="connected line 2", kind="lab_note", **kwargs):
    return new_annotation(kind=kind, text=text, timestamp=at(seconds),
                          origin=kwargs.pop("origin", "typed"),
                          recorded_at=kwargs.pop("recorded_at", at(seconds)), **kwargs)


# Model ---------------------------------------------------------------------


def test_note_needs_text_type_and_timezone():
    with pytest.raises(ValueError, match="text"):
        note(0, text="   ")
    with pytest.raises(ValueError, match="type"):
        note(0, kind="observation")
    with pytest.raises(ValueError, match="timezone"):
        new_annotation(kind="lab_note", text="x", timestamp=datetime(2026, 1, 1),
                       origin="typed")
    with pytest.raises(ValueError, match="at most"):
        note(0, text="x" * 2_001)


def test_edits_replace_and_deletes_remove_but_keep_creation_time():
    first = note(10, "bubble in line")
    second = note(5, "refilled", kind="qc_flag")
    original = resolve_annotations([first, second])
    assert [n.text for n in original] == ["refilled", "bubble in line"]

    bubble = original[1]
    moved = edited_annotation(bubble, timestamp=at(2), uncertainty_seconds=5.0,
                              origin="graph", signal=("bpr", "pressure"),
                              recorded_at=at(20))
    reworded = edited_annotation(bubble, text="bubble in line 2", recorded_at=at(21))
    removed = deleted_annotation(second.annotation_id, recorded_at=at(22))

    resolved = resolve_annotations([first, second, moved, removed])
    assert len(resolved) == 1
    assert resolved[0].timestamp == at(2)
    assert resolved[0].signal == ("bpr", "pressure")
    assert resolved[0].created_at == at(10)
    assert resolved[0].edited
    # Editing text alone keeps how and where the note was placed.
    assert reworded.timestamp == bubble.timestamp
    assert reworded.origin == "typed" and reworded.device_id is None


def test_corrections_to_unknown_or_deleted_notes_are_ignored():
    created = note(0)
    current = resolve_annotations([created])[0]
    records = [
        created,
        deleted_annotation(created.annotation_id),
        edited_annotation(current, text="too late"),
        deleted_annotation("missing"),
    ]
    assert resolve_annotations(records) == ()


@pytest.mark.parametrize(("raw", "expected"), [
    (0.3, 1), (1, 1), (4.6, 5), (7, 10), (45, 60), (61, 120), (400, 600),
    (5_000, 7_200), (10**7, 86_400),
])
def test_uncertainty_rounds_up_to_a_readable_step(raw, expected):
    assert nice_uncertainty(raw) == expected


def test_cell_text_shows_uncertainty_and_edits():
    created = note(0, "line\nreconnected", origin="graph", uncertainty_seconds=120.0)
    current = resolve_annotations([created])[0]
    assert format_annotation_cell(current) == "line reconnected (±2 min)"
    edited = resolve_annotations([created, edited_annotation(current, text="x")])[0]
    assert format_annotation_cell(edited) == "x (±2 min) [edited]"


def test_record_round_trips_through_json():
    record = note(3, origin="graph", uncertainty_seconds=5.0, signal=("p", "pressure"))
    assert annotation_record_from_dict(
        json.loads(json.dumps(annotation_record_to_dict(record)))
    ) == record


def test_journal_reader_leaves_a_partial_last_line(tmp_path):
    assert read_annotation_journal(tmp_path) == ((), 0)
    line = json.dumps(annotation_record_to_dict(note(0))) + "\n"
    (tmp_path / ANNOTATIONS_JOURNAL).write_bytes((line + line[:15]).encode())
    records, size = read_annotation_journal(tmp_path)
    assert len(records) == 1
    assert size == len(line.encode())


def test_newest_run_review_wins(tmp_path):
    assert latest_run_review(tmp_path) is None
    append_run_review(tmp_path, RunReview("minor", "leak at 40 min", at(100)))
    append_run_review(tmp_path, RunReview("severe", "leak from 30 min", at(200)))
    review = latest_run_review(tmp_path)
    assert review is not None
    assert (review.overall_qc, review.end_notes) == ("severe", "leak from 30 min")
    with pytest.raises(ValueError):
        RunReview("fine", "", at(0))


# Recording -----------------------------------------------------------------


def metadata() -> ExperimentMetadata:
    return ExperimentMetadata(experiment_id="EXP-1", operator="op",
                              started_at=START, notes="testing flow")


def test_recorder_refuses_notes_when_not_recording_instead_of_dropping():
    writer = InMemoryExperimentWriter()
    recorder = ExperimentRecorder(writer_factory=lambda _root: writer)
    with pytest.raises(RuntimeError, match="while recording"):
        recorder.record_annotation(note(0))
    recorder.start(metadata=metadata(), root_directory=".")
    assert recorder.recording_started_at == START
    recorder.record_annotation(note(1))
    assert len(writer.annotations) == 1
    recorder.stop()
    assert recorder.recording_started_at is None
    assert recorder.metadata.notes == "testing flow"


def test_directory_writer_journals_notes(tmp_path):
    writer = DirectoryExperimentWriter(tmp_path, live_export_interval_seconds=0)
    writer.open_experiment(metadata())
    writer.write_annotation(note(0, "first"))
    directory = writer.experiment_directory
    assert [n.text for n in read_annotations(directory)] == ["first"]
    writer.close_experiment()
    with pytest.raises(RuntimeError):
        writer.write_annotation(note(1))


# Exports -------------------------------------------------------------------


def prepare(directory):
    (directory / "metadata.json").write_text(
        json.dumps({"experiment_id": "test", "notes": "testing flow"}), encoding="utf-8")
    (directory / "measurements.journal.jsonl").touch()
    (directory / "events.journal.jsonl").touch()


def add_readings(directory, seconds):
    with (directory / "measurements.journal.jsonl").open("a", encoding="utf-8") as stream:
        for second in seconds:
            stream.write(json.dumps(dict(
                device_id="bpr", channel="pressure", unit="bar", value=1.0 + second,
                timestamp=at(second).isoformat(), quality="good")) + "\n")


def add_records(directory, *records):
    with (directory / ANNOTATIONS_JOURNAL).open("a", encoding="utf-8") as stream:
        for record in records:
            stream.write(json.dumps(annotation_record_to_dict(record)) + "\n")


def csv_rows(directory):
    with (directory / f"{directory.name}.csv").open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.reader(stream))


def test_csv_puts_notes_and_qc_flags_in_their_own_columns(tmp_path):
    prepare(tmp_path)
    add_readings(tmp_path, range(5))
    add_records(tmp_path, note(1.2, "line 2 on"), note(1, "bubble", kind="qc_flag"),
                note(0.9, "valve", kind="qc_flag"))
    IncrementalCsvExporter(tmp_path).update(final=True)
    rows = csv_rows(tmp_path)
    assert rows[0] == ["timestamp_utc", "bpr.pressure [bar]", "notes", "qc_flags"]
    assert rows[2][2:] == ["line 2 on", "valve | bubble"]
    assert rows[1][2:] == rows[3][2:] == ["", ""]


def test_notes_without_readings_get_their_own_rows(tmp_path):
    prepare(tmp_path)
    add_readings(tmp_path, [0, 10])
    add_records(tmp_path, note(-3, "before"), note(5, "between"), note(20, "after"))
    IncrementalCsvExporter(tmp_path).update(final=True)
    rows = csv_rows(tmp_path)
    assert [(row[1], row[2]) for row in rows[1:]] == [
        ("", "before"), ("1.0", ""), ("", "between"), ("11.0", ""), ("", "after"),
    ]
    headers, expected = build_wide_rows(
        [json.loads(line) for line in
         (tmp_path / "measurements.journal.jsonl").read_text().splitlines()],
        annotations=read_annotations(tmp_path),
    )
    assert rows == [headers, *[[str(cell) for cell in row] for row in expected]]


def test_note_for_an_earlier_time_rebuilds_rows_already_written(tmp_path):
    prepare(tmp_path)
    add_readings(tmp_path, range(10))
    exporter = IncrementalCsvExporter(tmp_path)
    exporter.update()
    assert len(csv_rows(tmp_path)) > 3
    created = note(2, "placed from graph", origin="graph", uncertainty_seconds=5.0)
    add_records(tmp_path, created)
    exporter.update()
    assert csv_rows(tmp_path)[3][2] == "placed from graph (±5 s)"

    current = read_annotations(tmp_path)[0]
    add_records(tmp_path, edited_annotation(current, timestamp=at(4), uncertainty_seconds=5.0,
                                            origin="graph"))
    exporter.update()
    rows = csv_rows(tmp_path)
    assert rows[3][2] == ""
    assert rows[5][2] == "placed from graph (±5 s) [edited]"

    add_records(tmp_path, deleted_annotation(current.annotation_id))
    exporter.update(final=True)
    assert all(row[2] == "" for row in csv_rows(tmp_path)[1:])


def test_excel_has_notes_sheet_and_review_on_overview(tmp_path):
    prepare(tmp_path)
    add_readings(tmp_path, range(3))
    add_records(tmp_path, note(1, "=SUM(A1)", kind="qc_flag", origin="graph",
                               uncertainty_seconds=5.0, signal=("bpr", "pressure")))
    append_run_review(tmp_path, RunReview("minor", "switched to testing leaks", at(60)))
    workbook = load_workbook(export_excel(tmp_path), read_only=True)
    try:
        assert workbook.sheetnames == ["Overview", "Data", "Events", "Notes"]
        overview = {row[0]: row[1] for row in workbook["Overview"].values if len(row) > 1}
        assert overview["Overall QC"] == "Minor"
        assert overview["Start notes"] == "testing flow"
        assert overview["End notes"] == "switched to testing leaks"
        assert overview["Notes during run"] == 1
        notes = list(workbook["Notes"].values)
        assert notes[0][:6] == ("timestamp", "type", "text", "uncertainty", "placed", "signal")
        assert notes[1][1:6] == ("QC flag", "=SUM(A1)", "5 s", "from graph", "bpr.pressure")
        data = list(workbook["Data"].values)
        assert data[2][-1] == "=SUM(A1) (±5 s)"
    finally:
        workbook.close()


def test_runs_recorded_before_notes_existed_still_export(tmp_path):
    prepare(tmp_path)
    add_readings(tmp_path, range(2))
    workbook = load_workbook(export_excel(tmp_path), read_only=True)
    try:
        overview = {row[0]: row[1] for row in workbook["Overview"].values if len(row) > 1}
        assert overview["Overall QC"] == "Not reviewed"
        assert len(list(workbook["Notes"].values)) == 1
    finally:
        workbook.close()
