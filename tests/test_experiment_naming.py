import json
from datetime import UTC, datetime
from functools import partial

import pytest
from openpyxl import load_workbook

from rig_control.control.service import RigControlService
from rig_control.data.directory_writer import DirectoryExperimentWriter
from rig_control.data.experiment_naming import (
    clean_id_part,
    existing_recording,
    make_experiment_id,
    remove_recording,
)
from rig_control.devices.manager import DeviceManager
from rig_control.devices.simulated_sensor import SimulatedSensor
from rig_control.experiment_recording import ExperimentRecorder
from rig_control.polling import PollingService
from rig_control.remembered_values import RememberedValues
from rig_control.ui.operation.model import OperationViewModel


STARTED = datetime(2026, 10, 22, 14, 9, 37).astimezone()


def test_id_is_initials_date_time_and_description() -> None:
    assert make_experiment_id("TPS", STARTED, "humidity tests") == (
        "TPS_2026-10-22_14h09_humidity-tests"
    )


def test_blank_description_ends_the_id_at_the_time() -> None:
    assert make_experiment_id("TPS", STARTED, "  ") == "TPS_2026-10-22_14h09"


def test_id_always_splits_into_its_parts() -> None:
    experiment_id = make_experiment_id(
        "T_P S", STARTED, 'flow 5 mL/min_test: "v2.1" <wet>'
    )

    assert experiment_id == "T-P-S_2026-10-22_14h09_flow-5-mL-min-test-v2-1-wet"
    assert len(experiment_id.split("_")) == 4


def test_long_description_is_capped_without_a_trailing_hyphen() -> None:
    description = make_experiment_id("TPS", STARTED, "a" * 39 + " b c").split("_")[3]

    assert description == "a" * 39
    assert clean_id_part("--x--") == "x"


def test_operator_initials_are_required() -> None:
    with pytest.raises(ValueError, match="initials"):
        make_experiment_id(" ", STARTED, "test")


def test_only_recording_folders_are_ever_removed(tmp_path) -> None:
    other = tmp_path / "TPS_2026-10-22_14h09"
    other.mkdir()
    (other / "keep.txt").write_text("not a recording")

    with pytest.raises(FileExistsError, match="not an experiment recording"):
        remove_recording(other)
    assert (other / "keep.txt").exists()


def make_model():
    manager = DeviceManager()
    manager.register(SimulatedSensor("probe", 25.0, "degC"))
    recorder = ExperimentRecorder(writer_factory=partial(
        DirectoryExperimentWriter, live_export_interval_seconds=0))
    polling = PollingService(manager, interval_seconds=60.0,
                             batch_handler=recorder.record_batch)
    model = OperationViewModel(manager, polling, recorder, RigControlService(manager),
                               profile_id="simulation")
    model.start_monitoring()
    return model


def record(model, tmp_path, **details):
    result = model.start_recording(
        operator="TPS", output_directory=tmp_path, started_at=STARTED, **details
    )
    if result.succeeded:
        model.stop_recording()
    return result


def test_recording_folder_and_exports_are_named_after_the_id(tmp_path) -> None:
    model = make_model()
    try:
        assert record(model, tmp_path, description="humidity tests",
                      electrolyser_id=" B142 ").succeeded
    finally:
        model.shutdown()

    folder = tmp_path / "TPS_2026-10-22_14h09_humidity-tests"
    metadata = json.loads((folder / "metadata.json").read_text(encoding="utf-8"))
    assert metadata["experiment_id"] == folder.name
    assert metadata["electrolyser_id"] == "B142"
    assert (folder / f"{folder.name}.csv").is_file()
    workbook = load_workbook(folder / f"{folder.name}.xlsx", read_only=True)
    try:
        overview = {row[0]: row[1] for row in workbook["Overview"].values if len(row) > 1}
    finally:
        workbook.close()
    assert overview["Experiment ID"] == folder.name
    assert overview["Electrolyser ID"] == "B142"


def test_blank_electrolyser_is_exported_as_not_recorded(tmp_path) -> None:
    model = make_model()
    try:
        assert record(model, tmp_path).succeeded
    finally:
        model.shutdown()

    folder = tmp_path / "TPS_2026-10-22_14h09"
    workbook = load_workbook(folder / f"{folder.name}.xlsx", read_only=True)
    try:
        overview = {row[0]: row[1] for row in workbook["Overview"].values if len(row) > 1}
    finally:
        workbook.close()
    assert overview["Electrolyser ID"] == "not recorded"


def test_existing_recording_is_kept_unless_replacement_is_confirmed(tmp_path) -> None:
    model = make_model()
    try:
        assert record(model, tmp_path, notes="first").succeeded
        folder = existing_recording(tmp_path, "TPS_2026-10-22_14h09")
        assert folder is not None

        refused = record(model, tmp_path, notes="second")
        assert not refused.succeeded and "already exists" in refused.summary
        kept = json.loads((folder / "metadata.json").read_text(encoding="utf-8"))
        assert kept["notes"] == "first"

        assert record(model, tmp_path, notes="second", replace_existing=True).succeeded
    finally:
        model.shutdown()
    replaced = json.loads((folder / "metadata.json").read_text(encoding="utf-8"))
    assert replaced["notes"] == "second"


def test_remembered_operator_survives_a_new_session(tmp_path) -> None:
    path = tmp_path / "settings" / "remembered.json"
    RememberedValues(path).set("operator", "TPS")

    assert RememberedValues(path).get("operator") == "TPS"


def test_unreadable_remembered_file_just_starts_empty(tmp_path) -> None:
    path = tmp_path / "remembered.json"
    path.write_text("{not json", encoding="utf-8")

    assert RememberedValues(path).get("operator") == ""
    RememberedValues(path).set("operator", "TPS")
    assert RememberedValues(path).get("operator") == "TPS"

